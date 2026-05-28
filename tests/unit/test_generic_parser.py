"""Tests for the format-agnostic content parsers."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from food_cpg_intelligence.data.generic_parser import (
    parse_any,
    parse_directory,
    parse_docx,
    parse_pptx,
    parse_xlsx,
)


def _make_docx(path: Path, title: str, body: str) -> None:
    doc = Document()
    title_para = doc.add_paragraph()
    title_run = title_para.add_run(title)
    title_run.bold = True
    doc.add_paragraph(body)
    doc.save(str(path))


def _make_xlsx(path: Path, rows: list[list[str]]) -> None:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(str(path))
    wb.close()


def _make_pptx(path: Path, slides: list[str]) -> None:
    pres = Presentation()
    layout = pres.slide_layouts[5]
    for text in slides:
        slide = pres.slides.add_slide(layout)
        if slide.shapes.title is not None:
            slide.shapes.title.text = text
    pres.save(str(path))


@pytest.mark.unit
def test_parse_docx_extracts_title_and_body(tmp_path: Path) -> None:
    src = tmp_path / "TestDoc.docx"
    _make_docx(src, title="A Catchy Title", body="A short body paragraph.")
    doc = parse_docx(src, content_type="membership", prefix="membership")
    assert doc.doc_id == "membership-TestDoc"
    assert doc.title == "A Catchy Title"
    assert "short body" in doc.full_text
    assert doc.content_type == "membership"
    assert doc.source_file == "TestDoc.docx"
    assert doc.metadata["format"] == "docx"
    assert doc.word_count > 0


@pytest.mark.unit
def test_parse_docx_handles_empty_file(tmp_path: Path) -> None:
    src = tmp_path / "Empty.docx"
    doc = Document()
    doc.save(str(src))
    parsed = parse_docx(src, content_type="reference", prefix="ref")
    assert parsed.word_count == 0
    assert parsed.full_text == ""
    assert parsed.doc_id == "ref-Empty"


@pytest.mark.unit
def test_parse_xlsx_flattens_rows_into_text(tmp_path: Path) -> None:
    src = tmp_path / "Prices.xlsx"
    _make_xlsx(src, [["SKU", "Price", "Notes"], ["A1", "1.99", "loss leader"], ["B2", "3.49", ""]])
    doc = parse_xlsx(src, content_type="reference", prefix="reference")
    assert "SKU | Price | Notes" in doc.full_text
    assert "A1 | 1.99 | loss leader" in doc.full_text
    assert doc.metadata["sheet_count"] == 1
    assert doc.metadata["row_count"] == 3
    assert doc.title == "Prices"


@pytest.mark.unit
def test_parse_pptx_extracts_slide_text(tmp_path: Path) -> None:
    src = tmp_path / "Deck.pptx"
    _make_pptx(src, ["Slide one heading", "Slide two heading", "Slide three heading"])
    doc = parse_pptx(src, content_type="reference", prefix="reference")
    assert doc.metadata["slide_count"] == 3
    assert "Slide one heading" in doc.full_text
    assert "Slide three heading" in doc.full_text


@pytest.mark.unit
def test_parse_any_dispatches_by_extension(tmp_path: Path) -> None:
    docx_path = tmp_path / "Doc.docx"
    _make_docx(docx_path, title="Doc Title", body="body")
    parsed = parse_any(docx_path, content_type="membership", prefix="m")
    assert parsed is not None
    assert parsed.title == "Doc Title"


@pytest.mark.unit
def test_parse_any_returns_none_for_unsupported(tmp_path: Path) -> None:
    txt_path = tmp_path / "notes.txt"
    txt_path.write_text("hello")
    assert parse_any(txt_path, content_type="reference", prefix="r") is None


@pytest.mark.unit
def test_parse_directory_skips_unsupported_and_empty(tmp_path: Path) -> None:
    _make_docx(tmp_path / "A.docx", title="A", body="body A is here")
    _make_docx(tmp_path / "B.docx", title="B", body="body B is here")
    (tmp_path / "ignore.txt").write_text("noise")

    docs = parse_directory(tmp_path, content_type="membership", prefix="membership")
    assert len(docs) == 2
    ids = {d.doc_id for d in docs}
    assert ids == {"membership-A", "membership-B"}


@pytest.mark.unit
def test_parse_directory_returns_empty_when_no_supported_files(tmp_path: Path) -> None:
    (tmp_path / "ignore.txt").write_text("noise")
    docs = parse_directory(tmp_path, content_type="reference", prefix="r")
    assert docs == []


@pytest.mark.unit
def test_doc_id_normalises_spaces(tmp_path: Path) -> None:
    src = tmp_path / "Spaced Out Name.docx"
    _make_docx(src, title="x", body="word word word")
    doc = parse_docx(src, content_type="reference", prefix="ref")
    assert doc.doc_id == "ref-Spaced_Out_Name"
